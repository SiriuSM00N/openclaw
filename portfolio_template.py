#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
美术作品集 PPT 模板生成器
风格：简约创意艺术风
"""

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

# 创建演示文稿
prs = Presentation()
prs.slide_width = Inches(13.333)  # 16:9 比例
prs.slide_height = Inches(7.5)

# 定义配色方案（参考简约艺术风格）
COLORS = {
    'bg': RGBColor(255, 255, 255),      # 白色背景
    'primary': RGBColor(45, 45, 45),     # 深灰文字
    'accent': RGBColor(100, 149, 237),   # 矢车菊蓝点缀
    'light_gray': RGBColor(240, 240, 240),
    'dark_gray': RGBColor(80, 80, 80),
}

def add_cover_slide(prs):
    """封面页"""
    slide = prs.slides.add_slide(prs.slide_layouts[6])  # 空白版式
    
    # 背景
    bg = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height
    )
    bg.fill.solid()
    bg.fill.fore_color.rgb = COLORS['bg']
    bg.line.fill.background()
    
    # 装饰色块（左侧）
    accent_bar = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, 0, 0, Inches(0.5), prs.slide_height
    )
    accent_bar.fill.solid()
    accent_bar.fill.fore_color.rgb = COLORS['accent']
    accent_bar.line.fill.background()
    
    # 标题
    title_box = slide.shapes.add_textbox(Inches(1.5), Inches(2.5), Inches(10), Inches(2))
    title_frame = title_box.text_frame
    title_para = title_frame.paragraphs[0]
    title_para.text = "我的作品集"
    title_para.font.size = Pt(54)
    title_para.font.bold = True
    title_para.font.color.rgb = COLORS['primary']
    title_para.font.name = "Microsoft YaHei"
    
    # 副标题
    subtitle_box = slide.shapes.add_textbox(Inches(1.5), Inches(4.2), Inches(10), Inches(1))
    subtitle_frame = subtitle_box.text_frame
    subtitle_para = subtitle_frame.paragraphs[0]
    subtitle_para.text = "Portfolio 2024-2025"
    subtitle_para.font.size = Pt(24)
    subtitle_para.font.color.rgb = COLORS['dark_gray']
    subtitle_para.font.name = "Microsoft YaHei"
    
    # 姓名
    name_box = slide.shapes.add_textbox(Inches(1.5), Inches(5.5), Inches(10), Inches(0.8))
    name_frame = name_box.text_frame
    name_para = name_frame.paragraphs[0]
    name_para.text = "小天狼星 | 游戏策划"
    name_para.font.size = Pt(18)
    name_para.font.color.rgb = COLORS['accent']
    name_para.font.name = "Microsoft YaHei"
    
    return slide

def add_about_slide(prs):
    """个人简介页"""
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    
    # 背景
    bg = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height
    )
    bg.fill.solid()
    bg.fill.fore_color.rgb = COLORS['bg']
    bg.line.fill.background()
    
    # 左侧装饰
    accent_bar = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, 0, 0, Inches(0.5), prs.slide_height
    )
    accent_bar.fill.solid()
    accent_bar.fill.fore_color.rgb = COLORS['accent']
    accent_bar.line.fill.background()
    
    # 页面标题
    title_box = slide.shapes.add_textbox(Inches(1.5), Inches(0.8), Inches(10), Inches(1))
    title_frame = title_box.text_frame
    title_para = title_frame.paragraphs[0]
    title_para.text = "关于我"
    title_para.font.size = Pt(36)
    title_para.font.bold = True
    title_para.font.color.rgb = COLORS['primary']
    title_para.font.name = "Microsoft YaHei"
    
    # 照片占位符（圆形）
    photo = slide.shapes.add_shape(
        MSO_SHAPE.OVAL, Inches(1.5), Inches(2), Inches(3), Inches(3)
    )
    photo.fill.solid()
    photo.fill.fore_color.rgb = COLORS['light_gray']
    photo.line.color.rgb = COLORS['accent']
    photo.line.width = Pt(2)
    
    # 照片提示文字
    photo_text = slide.shapes.add_textbox(Inches(1.5), Inches(3.2), Inches(3), Inches(0.5))
    photo_text_frame = photo_text.text_frame
    photo_text_para = photo_text_frame.paragraphs[0]
    photo_text_para.text = "照片"
    photo_text_para.font.size = Pt(14)
    photo_text_para.font.color.rgb = COLORS['dark_gray']
    photo_text_para.alignment = PP_ALIGN.CENTER
    
    # 简介文字
    about_box = slide.shapes.add_textbox(Inches(5.5), Inches(2), Inches(7), Inches(4))
    about_frame = about_box.text_frame
    about_frame.word_wrap = True
    
    about_lines = [
        ("姓名：小天狼星", 18),
        ("职业：游戏策划", 18),
        ("经验：5 年 + 手游经验", 18),
        ("", 10),
        ("擅长领域：", 18, True),
        ("• 系统设计", 16),
        ("• 数值平衡", 16),
        ("• 技术理解", 16),
        ("• 数据分析", 16),
    ]
    
    for i, line in enumerate(about_lines):
        para = about_frame.add_paragraph() if i > 0 else about_frame.paragraphs[0]
        para.text = line[0]
        para.font.size = Pt(line[1])
        para.font.color.rgb = COLORS['primary'] if len(line) < 3 else COLORS['accent']
        para.font.name = "Microsoft YaHei"
    
    return slide

def add_work_slide(prs, title, work_num):
    """作品展示页"""
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    
    # 背景
    bg = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height
    )
    bg.fill.solid()
    bg.fill.fore_color.rgb = COLORS['bg']
    bg.line.fill.background()
    
    # 左侧装饰
    accent_bar = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, 0, 0, Inches(0.5), prs.slide_height
    )
    accent_bar.fill.solid()
    accent_bar.fill.fore_color.rgb = COLORS['accent']
    accent_bar.line.fill.background()
    
    # 作品编号
    num_box = slide.shapes.add_textbox(Inches(1.5), Inches(0.5), Inches(2), Inches(0.8))
    num_frame = num_box.text_frame
    num_para = num_frame.paragraphs[0]
    num_para.text = f"WORK {work_num:02d}"
    num_para.font.size = Pt(14)
    num_para.font.color.rgb = COLORS['accent']
    num_para.font.name = "Microsoft YaHei"
    
    # 作品标题
    title_box = slide.shapes.add_textbox(Inches(1.5), Inches(0.5), Inches(10), Inches(1.2))
    title_frame = title_box.text_frame
    title_para = title_frame.paragraphs[0]
    title_para.text = title
    title_para.font.size = Pt(32)
    title_para.font.bold = True
    title_para.font.color.rgb = COLORS['primary']
    title_para.font.name = "Microsoft YaHei"
    title_para.alignment = PP_ALIGN.RIGHT
    
    # 作品图片占位符（大）
    img_placeholder = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(1.5), Inches(2), Inches(7), Inches(4)
    )
    img_placeholder.fill.solid()
    img_placeholder.fill.fore_color.rgb = COLORS['light_gray']
    img_placeholder.line.color.rgb = COLORS['accent']
    img_placeholder.line.width = Pt(2)
    img_placeholder.line.dash_style = 4  # 虚线
    
    # 图片提示
    img_text = slide.shapes.add_textbox(Inches(1.5), Inches(3.8), Inches(7), Inches(0.5))
    img_text_frame = img_text.text_frame
    img_text_para = img_text_frame.paragraphs[0]
    img_text_para.text = "作品截图 / 效果图"
    img_text_para.font.size = Pt(16)
    img_text_para.font.color.rgb = COLORS['dark_gray']
    img_text_para.alignment = PP_ALIGN.CENTER
    
    # 作品说明
    desc_box = slide.shapes.add_textbox(Inches(9.2), Inches(2), Inches(3.5), Inches(4))
    desc_frame = desc_box.text_frame
    desc_frame.word_wrap = True
    
    desc_lines = [
        ("项目概述", 18, True),
        ("在此处填写项目的", 14),
        ("简要描述和背景", 14),
        ("", 10),
        ("我的职责", 18, True),
        ("• 系统设计", 14),
        ("• 数值配置", 14),
        ("• 文档撰写", 14),
        ("", 10),
        ("成果数据", 18, True),
        ("• 上线时间", 14),
        ("• 用户数据", 14),
    ]
    
    for i, line in enumerate(desc_lines):
        para = desc_frame.add_paragraph() if i > 0 else desc_frame.paragraphs[0]
        para.text = line[0]
        para.font.size = Pt(line[1])
        para.font.color.rgb = COLORS['accent'] if len(line) > 2 else COLORS['primary']
        para.font.name = "Microsoft YaHei"
    
    return slide

def add_skill_slide(prs):
    """技能展示页"""
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    
    # 背景
    bg = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height
    )
    bg.fill.solid()
    bg.fill.fore_color.rgb = COLORS['bg']
    bg.line.fill.background()
    
    # 左侧装饰
    accent_bar = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, 0, 0, Inches(0.5), prs.slide_height
    )
    accent_bar.fill.solid()
    accent_bar.fill.fore_color.rgb = COLORS['accent']
    accent_bar.line.fill.background()
    
    # 页面标题
    title_box = slide.shapes.add_textbox(Inches(1.5), Inches(0.8), Inches(10), Inches(1))
    title_frame = title_box.text_frame
    title_para = title_frame.paragraphs[0]
    title_para.text = "专业技能"
    title_para.font.size = Pt(36)
    title_para.font.bold = True
    title_para.font.color.rgb = COLORS['primary']
    title_para.font.name = "Microsoft YaHei"
    
    # 技能卡片
    skills = [
        ("系统设计", "85%", 1.8),
        ("数值平衡", "80%", 3.2),
        ("技术理解", "75%", 4.6),
        ("数据分析", "70%", 6.0),
    ]
    
    for skill_name, skill_level, y_pos in skills:
        # 技能名称
        skill_box = slide.shapes.add_textbox(Inches(1.5), Inches(y_pos), Inches(3), Inches(0.5))
        skill_frame = skill_box.text_frame
        skill_para = skill_frame.paragraphs[0]
        skill_para.text = skill_name
        skill_para.font.size = Pt(18)
        skill_para.font.color.rgb = COLORS['primary']
        skill_para.font.name = "Microsoft YaHei"
        
        # 进度条背景
        bar_bg = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE, Inches(4.5), Inches(y_pos + 0.15), Inches(8), Inches(0.3)
        )
        bar_bg.fill.solid()
        bar_bg.fill.fore_color.rgb = COLORS['light_gray']
        bar_bg.line.fill.background()
        
        # 进度条填充
        level = float(skill_level.replace('%', '')) / 100
        bar_fill = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE, Inches(4.5), Inches(y_pos + 0.15), Inches(8 * level), Inches(0.3)
        )
        bar_fill.fill.solid()
        bar_fill.fill.fore_color.rgb = COLORS['accent']
        bar_fill.line.fill.background()
        
        # 百分比文字
        level_box = slide.shapes.add_textbox(Inches(12.6), Inches(y_pos), Inches(0.8), Inches(0.5))
        level_frame = level_box.text_frame
        level_para = level_frame.paragraphs[0]
        level_para.text = skill_level
        level_para.font.size = Pt(16)
        level_para.font.color.rgb = COLORS['accent']
        level_para.font.name = "Microsoft YaHei"
        level_para.alignment = PP_ALIGN.RIGHT
    
    return slide

def add_contact_slide(prs):
    """封底/联系方式页"""
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    
    # 背景
    bg = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height
    )
    bg.fill.solid()
    bg.fill.fore_color.rgb = COLORS['bg']
    bg.line.fill.background()
    
    # 左侧装饰
    accent_bar = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, 0, 0, Inches(0.5), prs.slide_height
    )
    accent_bar.fill.solid()
    accent_bar.fill.fore_color.rgb = COLORS['accent']
    accent_bar.line.fill.background()
    
    # 感谢文字
    thanks_box = slide.shapes.add_textbox(Inches(1.5), Inches(2), Inches(10), Inches(1.5))
    thanks_frame = thanks_box.text_frame
    thanks_para = thanks_frame.paragraphs[0]
    thanks_para.text = "感谢观看"
    thanks_para.font.size = Pt(48)
    thanks_para.font.bold = True
    thanks_para.font.color.rgb = COLORS['primary']
    thanks_para.font.name = "Microsoft YaHei"
    
    # 联系方式
    contact_box = slide.shapes.add_textbox(Inches(1.5), Inches(3.8), Inches(10), Inches(2.5))
    contact_frame = contact_box.text_frame
    contact_frame.word_wrap = True
    
    contact_lines = [
        ("联系方式", 24, True),
        ("", 10),
        ("📧 Email: your.email@example.com", 18),
        ("📱 微信：your_wechat", 18),
        ("💼 LinkedIn: your-linkedin", 18),
        ("", 10),
        ("期待与您合作！", 20, True),
    ]
    
    for i, line in enumerate(contact_lines):
        para = contact_frame.add_paragraph() if i > 0 else contact_frame.paragraphs[0]
        para.text = line[0]
        para.font.size = Pt(line[1])
        para.font.color.rgb = COLORS['accent'] if len(line) > 2 else COLORS['primary']
        para.font.name = "Microsoft YaHei"
    
    return slide

# 生成所有页面
print("🦐 正在生成美术作品集 PPT 模板...")

add_cover_slide(prs)      # 封面
add_about_slide(prs)      # 个人简介
add_work_slide(prs, "项目 A：XXX 游戏", 1)   # 作品 1
add_work_slide(prs, "项目 B：XXX 系统", 2)   # 作品 2
add_work_slide(prs, "项目 C：XXX 玩法", 3)   # 作品 3
add_skill_slide(prs)      # 技能展示
add_contact_slide(prs)    # 封底

# 保存文件
output_path = "/Users/sirius/.openclaw/workspace-default/美术作品集模板.pptx"
prs.save(output_path)

print(f"✅ PPT 模板已生成：{output_path}")
print(f"📄 共 7 页：封面 + 简介 + 3 个作品页 + 技能 + 封底")
print(f"🎨 风格：简约创意艺术风（白色背景 + 矢车菊蓝点缀）")
print(f"📝 提示：打开后替换占位文字和图片即可使用")
